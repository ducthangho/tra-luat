#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
normalize_document.py — Chuẩn hóa văn bản → doc record schema tra-luat-bundle/1.9 + cập nhật cache.
Bản tự đủ v1.9.14, tương thích CHÍNH XÁC interface mà ingest_to_cache.py::run_normalize gọi:
  --input --out --structure-level article --document-variant V --origin O --source-type S
  --cache --cache-dir D [--number-symbol] [--title] [--source-url] [--issuing-agency]
  [--effective-status] [--officialness-level]
Hỗ trợ input: .txt/.md (utf-8), .html (strip thẻ), .docx (python-docx nếu có, fallback unzip xml),
.pdf (pdftotext — KHÔNG convert OCR; trang scan sẽ thiếu và bị verify bắt), .doc (antiword nếu có).
Parse Điều chịu dính dòng (cùng regex verify_deep). Tự chạy T1 verify; coverage ghi trung thực.
"""
import argparse, hashlib, json, os, re, subprocess, sys, tempfile, unicodedata, html as _html
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from verify_deep import split_articles, t1_structure, ROMAN_PAT, extract_appendix, appendix_start  # noqa: E402
from keywords import extract_keywords  # noqa: E402  (v1.13.4: bắt tukhoa tất định lúc ingest)
from docx_markup import MODES, has_markup, strip_markup, resolve_mode, warn_line, scan_file  # noqa: E402  (r182)


def _clean_md_artifacts(t: str) -> str:
    """v1.15.2: làm sạch artifact của Google Docs read_file_content/markdown export
    (chỉ kích hoạt khi PHÁT HIỆN artifact ⇒ giữ nguyên hành vi với input đã sạch):
      - thực thể xuống dòng &#10; / &#xA; -> \\n  (để có whitespace TRƯỚC 'Điều N.')
      - bỏ backslash thoát markdown:  \\* \\. \\- \\[ \\] ...
      - bỏ chú thích/link footnote kiểu [\\[1\\]](url) và [1] sót lại
      - bỏ dấu nhấn ** / * (để ký tự trước 'Điều' là whitespace, không phải '*')
      - bỏ viền bảng markdown | và dòng phân tách |---|---|
    Nếu không vá, regex tách Điều (ART_SPLIT) trượt vì 'Điều' đứng sau '*' và không có newline."""
    has_artifacts = ("&#10;" in t) or ("\\*" in t) or ("](http" in t) \
        or ("**" in t) or bool(re.search(r"(?m)^\s*\|", t))
    if not has_artifacts:
        return t
    t = t.replace("&#10;", "\n").replace("&#xA;", "\n").replace("&#xa;", "\n")
    t = re.sub(r"\\([*_~`\[\]()#+.\-])", r"\1", t)            # unescape md punctuation
    t = re.sub(r"\[\s*\[?\d+\]?\s*\]\([^)]*\)", "", t)         # footnote link refs
    t = re.sub(r"(?<![\w])\[\d+\]", "", t)                     # [1] residues
    t = re.sub(r"\*{1,3}", "", t)                              # strip bold/italic asterisks
    t = re.sub(r"(?m)^\s*\|", "", t)                           # left table border
    t = re.sub(r"(?m)\|\s*$", "", t)                           # right table border
    t = re.sub(r"(?m)^\s*:?-{2,}:?\s*(?:\|\s*:?-{2,}:?\s*)*$", "", t)  # table separator rows
    t = t.replace(" | ", "  ")
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t


def _docx_zip_text(path: str, mode: str = "strip") -> tuple:
    """Nhánh KHÔNG python-docx: đọc thẳng word/document.xml. GIỮ bảng (w:tr→dòng, w:tc→tab).
    v1.35.12 (r182): lọc markup sửa đổi TRƯỚC khi bóc tag — nhánh này trước đây giữ CẢ w:delText
    (tệ hơn nhánh lib) ⇒ đọc chữ ĐÃ BỎ như quy phạm còn hiệu lực."""
    import zipfile
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8", "replace")
    xml, st = strip_markup(xml, mode)
    # GIỮ bảng: xuống dòng theo đoạn & hàng bảng, tab giữa ô
    xml = (xml.replace("</w:p>", "\n")
              .replace("</w:tr>", "\n")
              .replace("</w:tc>", "\t"))
    return _html.unescape(re.sub(r"<[^>]+>", "", xml)), st


def _docx_lib_text(path: str, mode: str = "strip") -> tuple:
    """Bóc text từ .docx GIỮ THỨ TỰ tài liệu, bao gồm CẢ bảng (w:tbl).

    Vá lỗi v1.15.6: trước đây dùng `docx.Document(path).paragraphs` — chỉ lấy
    đoạn văn, BỎ QUA bảng ⇒ phụ lục dạng bảng bị rớt khỏi text (vd Phụ lục IV
    'Danh mục ngành, nghề ĐTKD có điều kiện' của Luật Đầu tư). Hệ quả:
    `appendix_start`/`phu_luc` mất nguyên Phụ lục IV, đối chiếu sai (tưởng vắng
    ngành nghề). Nay duyệt body theo THỨ TỰ tài liệu, lấy cả w:p và w:tbl, để
    phụ lục bảng nằm đúng vị trí (sau thân) và `appendix_start` bắt được.

    Vá lỗi v1.35.12 (r182): `Paragraph(child, doc).text` nối MỌI run KHÔNG kiểm w:strike (giữ chữ ĐÃ
    BỎ) và python-docx BỎ run nằm trong w:ins/w:del (không phải con TRỰC TIẾP của w:p) ⇒ MẤT chữ MỚI
    THÊM. Bảng (c.text) CÙNG LỖI. Nay phần tử NÀO có markup thì lọc qua docx_markup rồi mới đọc —
    phần tử sạch giữ NGUYÊN đường cũ (bytes bất biến cho corpus đã nạp)."""
    import docx  # type: ignore
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from lxml import etree          # python-docx đã kéo lxml; tostring giữ NGUYÊN whitespace (khác
    #                                 `.xml` của python-docx = serialize_for_reading, có pretty-print)
    doc = docx.Document(path)
    tot = {}

    def _clean(el):
        """Phần tử có markup → lọc + tái dựng; sạch → trả nguyên (đường cũ, không đụng)."""
        x = etree.tostring(el, encoding="unicode")
        if not has_markup(x):
            return el
        x2, st = strip_markup(x, mode)
        for k, v in st.items():
            tot[k] = tot.get(k, 0) + v
        try:
            return parse_xml(x2)
        except Exception:  # noqa: BLE001 — xml lọc hỏng (hiếm) → giữ phần tử gốc, thà thừa hơn mất
            return el

    out: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            out.append(Paragraph(_clean(child), doc).text)
        elif tag == qn("w:tbl"):
            for row in Table(_clean(child), doc).rows:
                prev = None
                cells = []
                for c in row.cells:           # ô gộp (merge) trả lặp -> khử trùng liền kề
                    txt = c.text.strip()
                    if txt != prev:
                        cells.append(txt)
                    prev = txt
                out.append("\t".join(cells))
    return "\n".join(out), tot


def _docx_to_text(path: str, mode: str = "strip") -> str:
    """Điều phối 2 nhánh + CẢNH BÁO markup ra stderr (không nuốt: đọc nhầm dự thảo = sai pháp lý)."""
    try:
        import docx  # type: ignore  # noqa: F401
    except ImportError:
        txt, st = _docx_zip_text(path, mode)
    else:
        txt, st = _docx_lib_text(path, mode)
    w = warn_line(st, Path(path).name)
    if w:
        print(w, file=sys.stderr)
    return txt


def read_input(p: str, markup_mode: str = "strip") -> str:
    path = Path(p); ext = path.suffix.lower()
    if ext in (".txt", ".md"):
        return _clean_md_artifacts(path.read_text(encoding="utf-8", errors="replace"))
    if ext in (".html", ".htm"):
        h = path.read_text(encoding="utf-8", errors="replace")
        h = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", h)
        h = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</li>|</tr>", "\n", h)
        return _html.unescape(re.sub(r"(?s)<[^>]+>", " ", h))
    if ext == ".docx":
        return _docx_to_text(str(path), markup_mode)
    if ext == ".pdf":
        r = subprocess.run(["pdftotext", "-layout", str(path), "-"],
                           capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)
        if r.returncode != 0:
            raise RuntimeError("pdftotext lỗi: " + r.stderr[:300])
        return r.stdout
    if ext == ".doc":
        for cmd in (["antiword", str(path)], ["catdoc", str(path)]):
            try:
                r = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=60)
                if r.returncode == 0 and len(r.stdout) > 200:
                    return r.stdout
            except FileNotFoundError:
                continue
        raise RuntimeError(".doc cần antiword/catdoc, hoặc convert .doc→.docx trước")
    raise RuntimeError(f"Định dạng chưa hỗ trợ: {ext}")


def doc_key_of(so_hieu: str, variant: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", so_hieu.lower()).strip("-") + "--" + variant


def _write_json_atomic(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=path.name + ".", suffix=".tmp", dir=str(path.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as fh:
            json.dump(value, fh, ensure_ascii=False, indent=1, sort_keys=True)
            fh.write("\n")
        os.replace(tmp, path)
    except Exception:
        try:
            os.unlink(tmp)
        except FileNotFoundError:
            pass
        raise


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--structure-level", default="article")
    ap.add_argument("--document-variant", default="original")
    ap.add_argument("--origin", default="drive")
    ap.add_argument("--source-type", default="google_drive")
    ap.add_argument("--cache", action="store_true")
    ap.add_argument("--cache-dir", default=str(Path.home() / ".tra_luat_cache"))
    ap.add_argument("--number-symbol", default="")
    ap.add_argument("--title", default="")
    ap.add_argument("--source-url", default="")
    ap.add_argument("--issuing-agency", default="")
    ap.add_argument("--effective-status", default="")
    ap.add_argument("--officialness-level", default="official")
    ap.add_argument("--goc", default="", help="VBHN: số hiệu gốc + sđbs, phân tách ';'")
    ap.add_argument("--aliases", default="", help="Số hiệu tra cứu bổ sung, phân tách ';'")
    ap.add_argument("--as-appendix", action="store_true",
                    help="Văn bản CHỈ là phụ lục/mẫu: lưu nguyên khối phu_luc, không ép thành Điều, không FATAL.")
    ap.add_argument("--recorded-at", default="",
                    help="timestamp provenance tường minh; mặc định bỏ để cùng input tạo cùng cache bytes")
    ap.add_argument("--markup-mode", default="auto", choices=list(MODES),
                    help="(r182) markup sửa đổi .docx: auto=đánh dấu ⟦BỎ:…⟧ cho DỰ THẢO / loại hẳn cho "
                         "VBQPPL đã ban hành · strip=loại hẳn · mark=đánh dấu · raw=giữ nguyên (đối chứng)")
    a = ap.parse_args()

    # r182: dự thảo thì GIỮ dấu vết chữ bị đề nghị bỏ (người rà cần thấy); văn bản ban hành thì loại hẳn.
    _is_draft = ("draft" in a.document_variant.lower()
                 or "du_thao" in a.document_variant.lower()
                 or "dự thảo" in (a.effective_status or "").lower())
    _mk_mode = resolve_mode(a.markup_mode, is_draft=_is_draft)
    _mk_st = scan_file(a.input) if Path(a.input).suffix.lower() == ".docx" else {}
    _mk_st = {k: v for k, v in sorted(_mk_st.items()) if v}

    def _rd(_m):
        return unicodedata.normalize("NFC", read_input(a.input, _m).replace("\r\n", "\n").replace("\r", "\n"))

    # ★ CẤU TRÚC LUÔN PHÂN TÍCH TRÊN BẢN `strip` — ĐO ĐƯỢC (ca r182_gach_bo_khong_sinh_dieu_ma):
    # `ART_SPLIT` đòi CHỮ HOA ngay sau "Điều N." nên nhãn '⟦' CHẶN Điều thật, còn dấu ':' trong '⟦BỎ:'
    # lại THOẢ lookbehind [\s.;:] nên tiêu đề ĐÃ GẠCH BỎ đẻ ra "Điều ma". Không nhãn nào chữa được cả
    # hai chiều ⇒ bản `mark`/`raw` CHỈ là bản XEM (lưu ở text_markup), KHÔNG bao giờ nuôi parser.
    raw = _rd("strip")
    _raw_view = _rd(_mk_mode) if (_mk_st and _mk_mode != "strip") else ""
    if a.as_appendix:
        # v1.12.9: nạp văn bản PHỤ LỤC/MẪU độc lập — lưu nguyên khối, không yêu cầu Điều.
        body = ""           # không phân tích thân
        phu_luc = raw.strip()
        arts = split_articles(raw)   # giữ Điều nội bộ của mẫu (nếu có), KHÔNG bắt buộc
        rep = {"flags": ["APPENDIX: nạp dạng phụ lục/mẫu — không kiểm cấu trúc Điều nghị định"]}
        rep["t1"] = {"tier": "T1", "ok": True, "fatal": False, "articles_found": len(arts)}
        coverage = "appendix_only"
        a.document_variant = "appendix"
    else:
        # v1.12.8: TÁCH phụ lục TRƯỚC — phân tích cấu trúc chỉ trên phần THÂN (body),
        # tránh mẫu/biểu trong phụ lục bơm "Điều"/"Chương" giả gây T1 FATAL sai.
        _ap = appendix_start(raw)
        body = raw[:_ap] if _ap >= 0 else raw
        phu_luc = raw[_ap:].strip() if _ap >= 0 else ""
        arts = split_articles(body)
        if not arts:
            print(json.dumps({"error": "Không tách được Điều nào — kiểm tra nguồn "
                              "(nếu là phụ lục/mẫu độc lập, dùng --as-appendix)"}, ensure_ascii=False))
            sys.exit(2)
        # T1 nội bộ → coverage trung thực (chạy trên body, không gồm phụ lục)
        rep = {"flags": []}
        t1_structure(body, 0, a.number_symbol, rep)
        t1 = rep["t1"]
        if t1.get("fatal"):
            print(json.dumps({"error": "T1 FATAL: " + "; ".join(rep["flags"]),
                              "hint": "QUARANTINE — không nạp"}, ensure_ascii=False))
            sys.exit(2)
        coverage = "complete_vs_independent" if not rep["flags"] else (
            "partial: " + " | ".join(rep["flags"])[:400])

    # r182: markup sửa đổi là TÍN HIỆU PHÁP LÝ (bản đang sửa ≠ bản hiệu lực) — ghi vào record để
    # người tra thấy ngay. CHỈ ghi khi CÓ markup ⇒ văn bản sạch giữ NGUYÊN bytes cache như trước.
    if _mk_st:
        rep["flags"].append("MARKUP_SUA_DOI: " + ", ".join(f"{k}={v}" for k, v in _mk_st.items())
                            + f" (che_do={_mk_mode}) — chữ w:del/gạch-ngang ĐÃ BỎ, w:ins là chữ MỚI THÊM")

    nums = sorted(arts)
    norm = re.sub(r"[ \t]+", " ", raw).strip()
    so_hieu = a.number_symbol or "UNKNOWN"
    key = doc_key_of(so_hieu, a.document_variant)
    record = {
        "schema": "tra-luat-bundle/1.9",
        "doc_key": key, "so_hieu": so_hieu, "number_symbol": so_hieu,
        "title": a.title, "variant": a.document_variant,
        "goc": a.goc, "aliases": [s for s in a.aliases.split(";") if s],
        "issuing_agency": a.issuing_agency, "document_status": a.effective_status,
        "provenance": a.officialness_level, "officialness": a.officialness_level,
        "source": {"origin": a.origin, "source_type": a.source_type, "url": a.source_url,
                   "input_file": str(Path(a.input).name)},
        "text_hash": hashlib.sha256(norm.encode()).hexdigest(), "chars": len(raw),
        "structure": {"chapters": len(ROMAN_PAT.findall(body)),
                      "articles": len(nums), "terminal_article": (nums[-1] if nums else 0),
                      "has_phu_luc": bool(phu_luc), "phu_luc_chars": len(phu_luc),
                      "is_appendix": bool(getattr(a, "as_appendix", False))},
        "coverage": coverage, "verify_flags": rep["flags"],
        "articles": {str(n): arts[n] for n in nums},
        "has_phu_luc": bool(phu_luc),
        "phu_luc": phu_luc,
    }
    if _mk_st:
        record["markup"] = _mk_st
        record["markup_mode"] = _mk_mode
        if _raw_view:
            record["text_markup"] = _raw_view   # bản XEM (⟦BỎ:…⟧) — KHÔNG dùng để phân tích cấu trúc
    if a.recorded_at:
        record["verified_at"] = a.recorded_at
        record["ingested_at"] = a.recorded_at
    # v1.13.4: MÓC CHỦ ĐỀ tất định (Chương/Điều) — KHÔNG AI, idempotent. Để Claude lọc-thô
    # chọn nguồn trước khi notebook_query. Danh tính (so_hieu/loai/linhvuc) đã là cột riêng.
    record["tukhoa"] = extract_keywords(record)
    _write_json_atomic(Path(a.out), record)

    if a.cache:
        cd = Path(a.cache_dir); (cd / "docs").mkdir(parents=True, exist_ok=True)
        _write_json_atomic(cd / "docs" / f"{key}.json", record)
        bi_p = cd / "bundle_index.json"
        bi = json.loads(bi_p.read_text(encoding="utf-8")) if bi_p.exists() else \
            {"schema": "tra-luat-bundle/1.9", "docs": []}
        bi["docs"] = [d for d in bi["docs"] if d.get("doc_key") != key]
        _idx = {k: record[k] for k in ("doc_key", "so_hieu", "title", "variant", "text_hash", "coverage")}
        _idx["aliases"] = record["aliases"]
        if a.recorded_at:
            _idx["verified_at"] = a.recorded_at
        bi["docs"].append(_idx)
        bi["docs"] = sorted(bi["docs"], key=lambda d: str(d.get("doc_key") or ""))
        if a.recorded_at:
            bi["generatedAt"] = a.recorded_at
        else:
            bi.pop("generatedAt", None)
        _write_json_atomic(bi_p, bi)
    print(json.dumps({"ok": True, "doc_key": key, "articles": len(nums),
                      "coverage": coverage}, ensure_ascii=False))


if __name__ == "__main__":
    main()
